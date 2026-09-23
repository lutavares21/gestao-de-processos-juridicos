# -*- coding: utf-8 -*-
"""
App Flask - Gestão de Processos Jurídicos
"""

import csv
import io
import os
import re
import unicodedata
from datetime import datetime, date
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, abort, jsonify
from flask_login import (
    LoginManager, login_user, logout_user, login_required, current_user,
)
from sqlalchemy.exc import IntegrityError
from sqlalchemy import func, case
from werkzeug.security import check_password_hash, generate_password_hash
from models import (
    db, Processo, Parte, Advogado, Movimento, PedidoTrabalhista, RateioCR,
    PedidoCivel, TituloRecuperacao, Operador, RegistroAtividade,
)

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///processos.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
# Chave usada para proteger a sessão de login. Em algum momento vale trocar
# por uma variável de ambiente, mas por enquanto um valor fixo já funciona.
app.config["SECRET_KEY"] = "troque-esta-chave-por-uma-string-aleatoria-depois"

db.init_app(app)

with app.app_context():
    db.create_all()

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = "Faça login para acessar o sistema."


@login_manager.user_loader
def carregar_operador(operador_id):
    return db.session.get(Operador, int(operador_id))


@app.route("/logout-beacon", methods=["POST"])
def logout_beacon():
    """Chamada pelo navegador (via navigator.sendBeacon) só quando a página
    detectou que está sendo fechada de verdade - o próprio JavaScript
    (static/js/sessao.js) já filtrou navegações internas, F5, etc, então
    aqui é só deslogar direto, sem prazo de espera."""
    if current_user.is_authenticated:
        logout_user()
    return ("", 204)


def admin_required(f):
    """Só deixa passar se o operador logado for administrador.
    Quem não for recebe 403 (acesso negado)."""
    @wraps(f)
    @login_required
    def decorado(*args, **kwargs):
        if not current_user.eh_administrador:
            abort(403)
        return f(*args, **kwargs)
    return decorado


def registrar_atividade(acao, descricao):
    """Grava uma linha no log de atividades, associada ao operador
    logado no momento. Não faz commit sozinha - a chamada que já vai
    salvar o processo/operador salva essa linha junto."""
    db.session.add(RegistroAtividade(
        operador_id=current_user.id if current_user.is_authenticated else None,
        operador_nome=current_user.nome if current_user.is_authenticated else "Desconhecido",
        acao=acao,
        descricao=descricao,
    ))



@app.template_filter("moeda")
def formatar_moeda(valor):
    """Formata um número no padrão contábil brasileiro: R$ 1.234,56.
    Retorna '—' quando o valor é None."""
    if valor is None:
        return "—"
    texto = f"{float(valor):,.2f}"
    texto = texto.replace(",", "_").replace(".", ",").replace("_", ".")
    return f"R$ {texto}"


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    if request.method == "GET":
        return render_template("login.html")

    login_informado = request.form.get("login", "").strip()
    senha = request.form.get("senha", "")

    operador = Operador.query.filter(
        func.lower(Operador.login) == login_informado.lower()
    ).first()

    if operador and check_password_hash(operador.senha_hash, senha):
        login_user(operador)
        proxima = request.args.get("next")
        return redirect(proxima or url_for("index"))

    return render_template("login.html", erro_login=True)


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/operadores")
@admin_required
def operadores():
    lista = Operador.query.order_by(Operador.nome).all()
    return render_template("operadores.html", operadores=lista)


@app.route("/operadores/novo", methods=["GET", "POST"])
@admin_required
def operador_novo():
    if request.method == "GET":
        return render_template("novo_operador.html")

    form = request.form
    nome = form.get("nome", "").strip()
    login_novo = form.get("login", "").strip()
    email = form.get("email", "").strip()
    senha = form.get("senha", "")
    confirmar_senha = form.get("confirmar_senha", "")
    nivel = form.get("nivel", "comum")
    if nivel not in ("administrador", "comum"):
        nivel = "comum"

    erros = []
    if not nome:
        erros.append("Informe o nome.")
    if not login_novo:
        erros.append("Informe o login.")
    if not email:
        erros.append("Informe o e-mail.")
    if not senha:
        erros.append("Informe a senha.")
    if senha != confirmar_senha:
        erros.append("A senha e a confirmação de senha não coincidem.")
    if login_novo and Operador.query.filter(func.lower(Operador.login) == login_novo.lower()).first():
        erros.append("Já existe um operador com esse login.")

    if erros:
        return render_template("novo_operador.html", erros=erros, valores=form)

    novo_operador = Operador(
        nome=nome,
        login=login_novo,
        email=email,
        senha_hash=generate_password_hash(senha),
        nivel=nivel,
    )
    db.session.add(novo_operador)
    registrar_atividade("operador_criado", f"Criou o operador {nome} ({login_novo})")
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        erros.append("Já existe um operador com esse login.")
        return render_template("novo_operador.html", erros=erros, valores=form)

    return redirect(url_for("operadores"))


@app.route("/operadores/<int:operador_id>/editar", methods=["GET", "POST"])
@admin_required
def operador_editar(operador_id):
    operador = db.session.get(Operador, operador_id)
    if operador is None:
        abort(404)

    if request.method == "GET":
        return render_template("operador_editar.html", operador=operador)

    form = request.form
    nome = form.get("nome", "").strip()
    login_novo = form.get("login", "").strip()
    email = form.get("email", "").strip()
    senha = form.get("senha", "")
    confirmar_senha = form.get("confirmar_senha", "")
    nivel = form.get("nivel", "comum")
    if nivel not in ("administrador", "comum"):
        nivel = "comum"

    erros = []
    if not nome:
        erros.append("Informe o nome.")
    if not login_novo:
        erros.append("Informe o login.")
    if not email:
        erros.append("Informe o e-mail.")
    # A senha é opcional na edição - só troca se o campo for preenchido.
    if senha and senha != confirmar_senha:
        erros.append("A senha e a confirmação de senha não coincidem.")

    login_em_uso = Operador.query.filter(
        func.lower(Operador.login) == login_novo.lower(),
        Operador.id != operador.id,
    ).first()
    if login_novo and login_em_uso:
        erros.append("Já existe outro operador com esse login.")

    if erros:
        return render_template("operador_editar.html", operador=operador, erros=erros, valores=form)

    operador.nome = nome
    operador.login = login_novo
    operador.email = email
    operador.nivel = nivel
    if senha:
        operador.senha_hash = generate_password_hash(senha)

    try:
        registrar_atividade("operador_editado", f"Editou o operador {operador.nome} ({operador.login})")
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        erros.append("Já existe outro operador com esse login.")
        return render_template("operador_editar.html", operador=operador, erros=erros, valores=form)

    return redirect(url_for("operadores"))


@app.route("/operadores/<int:operador_id>/excluir", methods=["POST"])
@admin_required
def operador_excluir(operador_id):
    operador = db.session.get(Operador, operador_id)
    if operador is None:
        abort(404)

    # Não deixa o administrador excluir a própria conta enquanto está
    # logado com ela - evitaria ele mesmo se trancar fora do sistema.
    if operador.id == current_user.id:
        erros = ["Você não pode excluir o próprio usuário enquanto estiver logado com ele."]
        return render_template("operador_editar.html", operador=operador, erros=erros)

    registrar_atividade("operador_excluido", f"Excluiu o operador {operador.nome} ({operador.login})")
    db.session.delete(operador)
    db.session.commit()
    return redirect(url_for("operadores"))


@app.route("/atividades")
@admin_required
def atividades():
    registros = (
        RegistroAtividade.query
        .order_by(RegistroAtividade.data_hora.desc())
        .limit(300)
        .all()
    )
    return render_template("atividades.html", registros=registros)


def texto_para_data(valor):
    """Converte string 'AAAA-MM-DD' do formulário em objeto date. Retorna None se vazio."""
    if not valor:
        return None
    return datetime.strptime(valor, "%Y-%m-%d").date()


def texto_para_numero(valor):
    """Converte string do formulário em número decimal. Retorna None se vazio."""
    if not valor:
        return None
    return float(valor)


def preencher_campos_processo(processo, form):
    """Preenche (ou atualiza) os campos comuns do processo a partir do
    formulário. Usado tanto no cadastro novo quanto na edição."""
    processo.numero_processo = form.get("numero_processo")
    processo.juizado = form.get("juizado")
    processo.comarca = form.get("comarca")
    processo.uf = form.get("uf")
    processo.tipo_acao = form.get("tipo_acao")
    processo.data_distribuicao = texto_para_data(form.get("data_distribuicao"))
    processo.valor_causa = texto_para_numero(form.get("valor_causa"))
    processo.data_audiencia_1 = texto_para_data(form.get("data_audiencia_1"))
    processo.audiencia_1_horario = form.get("audiencia_1_horario") or None
    processo.audiencia_1_tipo = form.get("audiencia_1_tipo") or None
    processo.audiencia_1_link = form.get("audiencia_1_link") or None
    processo.data_audiencia_2 = texto_para_data(form.get("data_audiencia_2"))
    processo.data_audiencia_3 = texto_para_data(form.get("data_audiencia_3"))
    processo.data_arquivamento = texto_para_data(form.get("data_arquivamento"))
    processo.centro_resultado = form.get("centro_resultado")
    processo.escritorio = form.get("escritorio")
    processo.resultado = form.get("resultado")
    processo.sentenca = form.get("sentenca")
    processo.status = form.get("status", "ativo")
    processo.risco = form.get("risco")
    processo.grau_instancia = form.get("grau_instancia")
    processo.resumo = form.get("resumo")
    processo.valor_final = texto_para_numero(form.get("valor_final"))
    processo.honorarios_advogado = texto_para_numero(form.get("honorarios_advogado"))
    processo.honorarios_periciais = texto_para_numero(form.get("honorarios_periciais"))
    processo.custas_processuais = texto_para_numero(form.get("custas_processuais"))
    processo.deposito_recursal = texto_para_numero(form.get("deposito_recursal"))
    processo.valor_alvara = texto_para_numero(form.get("valor_alvara"))
    processo.data_alvara = texto_para_data(form.get("data_alvara"))
    processo.economia_gerada = texto_para_numero(form.get("economia_gerada"))


def montar_processo_base(form):
    """Cria um objeto Processo novo com os campos comuns já preenchidos."""
    processo = Processo()
    preencher_campos_processo(processo, form)
    return processo


def preencher_partes_e_advogados(processo, form, incluir_partes=True):
    """Preenche partes, advogados, movimentos e rateio - comum a todas as páginas.

    O rótulo salvo em Parte.tipo depende da origem do cadastro: processos
    trabalhistas usam 'reclamante'/'reclamada'; os demais (cível) usam
    'autor'/'reu'.

    incluir_partes=False é usado na edição: o reclamante/autor e a
    reclamada/réu não podem ser alterados depois do cadastro, então as
    partes existentes do processo são preservadas e o conteúdo enviado
    pelo formulário para autor_nome/reu_nome é ignorado."""
    if incluir_partes:
        if processo.origem_cadastro == "trabalhista":
            tipo_polo_ativo, tipo_polo_passivo = "reclamante", "reclamada"
        else:
            tipo_polo_ativo, tipo_polo_passivo = "autor", "reu"

        for nome in form.getlist("autor_nome"):
            if nome.strip():
                processo.partes.append(Parte(tipo=tipo_polo_ativo, nome=nome.strip()))
        for nome in form.getlist("reu_nome"):
            if nome.strip():
                processo.partes.append(Parte(tipo=tipo_polo_passivo, nome=nome.strip()))

    nomes_adv_autor = form.getlist("advogado_autor_nome")
    oabs_adv_autor = form.getlist("advogado_autor_oab")
    for nome, oab in zip(nomes_adv_autor, oabs_adv_autor):
        if nome.strip():
            processo.advogados.append(Advogado(lado="autor", nome=nome.strip(), oab=oab.strip()))

    nomes_adv_reu = form.getlist("advogado_reu_nome")
    oabs_adv_reu = form.getlist("advogado_reu_oab")
    for nome, oab in zip(nomes_adv_reu, oabs_adv_reu):
        if nome.strip():
            processo.advogados.append(Advogado(lado="reu", nome=nome.strip(), oab=oab.strip()))

    datas_mov = form.getlist("movimento_data")
    tipos_mov = form.getlist("movimento_tipo")
    for data_str, tipo in zip(datas_mov, tipos_mov):
        if tipo.strip():
            processo.movimentos.append(
                Movimento(data_movimento=texto_para_data(data_str), tipo_movimento=tipo.strip())
            )

    if form.get("centro_resultado") == "rateio":
        crs_marcados = form.getlist("rateio_cr")
        crs_digitados = form.getlist("rateio_cr_manual")
        for cr in crs_marcados + crs_digitados:
            if cr.strip():
                processo.rateio_crs.append(RateioCR(centro_resultado=cr.strip()))


@app.context_processor
def injetar_contadores_menu():
    """Disponibiliza as contagens de processos ativos (cível e trabalhista)
    para o menu lateral, em todas as páginas, sem precisar passar isso
    manualmente em cada rota."""
    return dict(
        contagem_civel_ativos=Processo.query.filter_by(origem_cadastro="civel", status="ativo").count(),
        contagem_trabalhista_ativos=Processo.query.filter_by(origem_cadastro="trabalhista", status="ativo").count(),
        contagem_recuperacao_ativos=Processo.query.filter_by(origem_cadastro="civel_recuperacao", status="ativo").count(),
    )


@app.route("/")
@login_required
def index():
    return render_template("inicio.html")


def numero_processo_ja_existe(numero_processo):
    """Verifica se já existe algum processo (cível ou trabalhista) cadastrado
    com esse número. Usado para bloquear cadastros duplicados."""
    return Processo.query.filter_by(numero_processo=numero_processo).first() is not None


def campos_obrigatorios_faltando(form, trabalhista=False):
    """Confere os campos que nenhum cadastro pode ficar sem: número do
    processo e pelo menos um nome de cada lado. Devolve a lista dos que
    faltam (vazia = pode salvar).

    O formulário sempre manda os nomes em autor_nome/reu_nome, mesmo no
    trabalhista - só os rótulos mostrados ao operador mudam.

    Essa checagem repete a do JavaScript de propósito: o navegador pode
    estar com script desativado, ou o POST pode vir de fora da tela."""
    rotulo_autor = "Reclamante" if trabalhista else "Autor"
    rotulo_reu = "Reclamada" if trabalhista else "Réu"

    faltando = []
    if not (form.get("numero_processo") or "").strip():
        faltando.append("Número do processo")
    if not any(nome.strip() for nome in form.getlist("autor_nome")):
        faltando.append(rotulo_autor)
    if not any(nome.strip() for nome in form.getlist("reu_nome")):
        faltando.append(rotulo_reu)
    return faltando


@app.route("/inicio")
@login_required
def inicio():
    return render_template("inicio.html")


@app.route("/cadastro/civel", methods=["GET", "POST"])
@login_required
def cadastro_civel():
    if request.method == "GET":
        sucesso = request.args.get("sucesso") == "1"
        return render_template("cadastro_civel.html", sucesso=sucesso)

    form = request.form
    numero_processo = form.get("numero_processo")

    faltando = campos_obrigatorios_faltando(form)
    if faltando:
        return render_template("cadastro_civel.html", erro_campos_obrigatorios=faltando)

    if numero_processo_ja_existe(numero_processo):
        return render_template(
            "cadastro_civel.html",
            erro_numero_duplicado=numero_processo,
        )

    processo = montar_processo_base(form)
    processo.origem_cadastro = "civel"
    preencher_partes_e_advogados(processo, form)

    # Pedidos e requerimentos (cível) - checkboxes marcados
    for descricao in form.getlist("pedido_civel"):
        if descricao.strip():
            processo.pedidos_civeis.append(PedidoCivel(descricao=descricao.strip()))

    db.session.add(processo)
    registrar_atividade("processo_criado", f"Cadastrou o processo cível {numero_processo}")
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return render_template("cadastro_civel.html", erro_numero_duplicado=numero_processo)

    return redirect(url_for("cadastro_civel", sucesso=1))


# ---------------------------------------------------------------------------
# Recuperação de Crédito - leitura da planilha de títulos
# ---------------------------------------------------------------------------
#
# A planilha é lida por uma rota separada (/recuperacao/ler-planilha), chamada
# por JavaScript. Ela devolve as linhas em JSON e o navegador monta a tabela
# editável na tela. Assim o operador pode corrigir qualquer célula antes de
# salvar, e a página do cadastro não recarrega (não se perde o que já foi
# digitado nas outras seções).


class PlanilhaErro(Exception):
    """Erro de leitura que já tem uma mensagem pronta para mostrar ao operador."""


# Nomes de coluna aceitos para cada campo. A comparação ignora acentos,
# maiúsculas e espaços sobrando, então "Correção", "correcao" e "CORREÇÃO "
# caem todos no mesmo lugar.
COLUNAS_RECUPERACAO = {
    "company": ["company", "companhia", "empresa", "filial"],
    "cliente": ["cliente", "sacado", "devedor", "razao social", "id cliente"],
    "nota_fiscal": ["nota fiscal", "notafiscal", "nf", "nfe", "nf-e",
                    "numero da nota", "numero nf", "no nf", "documento", "titulo"],
    "emissao": ["emissao", "data emissao", "data de emissao", "dt emissao"],
    "vencimento": ["vencimento", "data vencimento", "data de vencimento", "dt vencimento"],
    "valor": ["valor", "valor original", "valor da nota", "valor nf", "valor principal"],
    "saldo": ["saldo", "saldo devedor", "saldo em aberto"],
    "juros": ["juros", "juros de mora", "mora"],
    "correcao": ["correcao", "correcao monetaria", "atualizacao", "atualizacao monetaria"],
    "outros": ["outros", "outras despesas", "despesas", "acrescimos"],
    "total": ["total", "valor total", "total geral", "total atualizado"],
}

CAMPOS_NUMERICOS_RECUPERACAO = ["valor", "saldo", "juros", "correcao", "outros", "total"]
CAMPOS_DATA_RECUPERACAO = ["emissao", "vencimento"]

EXTENSOES_PLANILHA = (
    ".xlsx", ".xlsm", ".xltx", ".xltm",   # Excel moderno (openpyxl)
    ".xls",                                # Excel antigo (xlrd)
    ".csv", ".txt", ".tsv",                # texto separado (módulo csv)
    ".ods",                                # LibreOffice / OpenOffice (odfpy)
)


def normalizar_rotulo(texto):
    """Deixa um nome de coluna comparável: sem acento, minúsculo, sem
    pontuação e sem espaços repetidos."""
    if texto is None:
        return ""
    texto = str(texto)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.lower().replace("º", "").replace("°", "")
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return texto.strip()


def valor_para_numero_planilha(valor):
    """Converte o conteúdo de uma célula em número (float) ou None.

    Aceita número puro vindo do Excel, e também texto em formato brasileiro
    ('1.234,56', 'R$ 1.234,56', '(1.234,56)' para negativo) ou americano
    ('1,234.56')."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)

    texto = str(valor).strip()
    if not texto:
        return None

    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()")
    texto = re.sub(r"[^\d,.\-]", "", texto)  # tira "R$", espaços, etc.
    if not texto or texto in {"-", ".", ","}:
        return None

    if "," in texto and "." in texto:
        # O separador decimal é o que aparece por último.
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    else:
        # Só pontos: se houver mais de um, ou se o último grupo tiver 3
        # dígitos, é separador de milhar (1.234 / 1.234.567).
        partes = texto.split(".")
        if len(partes) > 2 or (len(partes) == 2 and len(partes[1]) == 3):
            texto = "".join(partes)

    try:
        numero = float(texto)
    except ValueError:
        return None
    return -numero if negativo else numero


def valor_para_data_planilha(valor):
    """Converte o conteúdo de uma célula em date, ou None se não der.

    Aceita datetime/date vindos do Excel, texto em vários formatos e também
    o número serial de data do Excel (ex.: 45000)."""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor

    # Número serial do Excel (contado a partir de 30/12/1899).
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            from datetime import timedelta
            return (datetime(1899, 12, 30) + timedelta(days=float(valor))).date()
        except (ValueError, OverflowError):
            return None

    texto = str(valor).strip()
    if not texto:
        return None
    texto = texto.split(" ")[0].split("T")[0]  # descarta a hora, se vier junto

    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y",
                    "%d.%m.%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _linhas_xlsx(stream):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise PlanilhaErro(
            "A biblioteca openpyxl não está instalada no servidor. "
            "Rode: pip3.10 install --user openpyxl"
        )
    # data_only=True traz o resultado das fórmulas, não a fórmula em si.
    planilha = load_workbook(stream, data_only=True, read_only=True)
    aba = planilha.active
    return [list(linha) for linha in aba.iter_rows(values_only=True)]


def _linhas_xls(conteudo):
    try:
        import xlrd
    except ImportError:
        raise PlanilhaErro(
            "Este é um Excel antigo (.xls) e a biblioteca xlrd não está "
            "instalada no servidor. Rode: pip3.10 install --user xlrd  "
            "(ou salve a planilha como .xlsx e anexe de novo)."
        )
    livro = xlrd.open_workbook(file_contents=conteudo)
    aba = livro.sheet_by_index(0)
    linhas = []
    for indice in range(aba.nrows):
        celulas = []
        for celula in aba.row(indice):
            if celula.ctype == xlrd.XL_CELL_DATE:
                ano, mes, dia, *_ = xlrd.xldate_as_tuple(celula.value, livro.datemode)
                celulas.append(date(ano, mes, dia) if ano else None)
            else:
                celulas.append(celula.value)
        linhas.append(celulas)
    return linhas


def _linhas_csv(conteudo):
    texto = None
    for codificacao in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:
        raise PlanilhaErro("Não foi possível identificar a codificação do arquivo de texto.")

    amostra = texto[:4000]
    try:
        separador = csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
    except csv.Error:
        # Chute razoável: no Brasil o Excel exporta CSV com ponto e vírgula.
        separador = ";" if amostra.count(";") >= amostra.count(",") else ","
    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=separador)]


def _linhas_ods(stream):
    try:
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf.text import P
    except ImportError:
        raise PlanilhaErro(
            "Este arquivo é .ods e a biblioteca odfpy não está instalada no "
            "servidor. Rode: pip3.10 install --user odfpy  (ou salve a "
            "planilha como .xlsx e anexe de novo)."
        )
    documento = load(stream)
    tabelas = documento.spreadsheet.getElementsByType(Table)
    if not tabelas:
        raise PlanilhaErro("O arquivo .ods não tem nenhuma aba com dados.")

    linhas = []
    for linha_ods in tabelas[0].getElementsByType(TableRow):
        celulas = []
        for celula in linha_ods.getElementsByType(TableCell):
            repeticoes = int(celula.getAttribute("numbercolumnsrepeated") or 1)
            bruto = (celula.getAttribute("value")
                     or celula.getAttribute("datevalue")
                     or "".join(str(paragrafo) for paragrafo in celula.getElementsByType(P)))
            celulas.extend([bruto or None] * min(repeticoes, 60))
        linhas.append(celulas)
    return linhas


def ler_linhas_brutas(arquivo):
    """Abre o arquivo enviado e devolve uma lista de listas com o conteúdo
    cru das células, escolhendo o leitor certo pela extensão."""
    nome = (arquivo.filename or "").strip()
    extensao = os.path.splitext(nome)[1].lower()

    if extensao not in EXTENSOES_PLANILHA:
        raise PlanilhaErro(
            f"Formato não suportado ({extensao or 'sem extensão'}). "
            "Aceitos: " + ", ".join(EXTENSOES_PLANILHA)
        )

    conteudo = arquivo.read()
    if not conteudo:
        raise PlanilhaErro("O arquivo enviado está vazio.")

    if extensao in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _linhas_xlsx(io.BytesIO(conteudo))
    if extensao == ".xls":
        return _linhas_xls(conteudo)
    if extensao == ".ods":
        return _linhas_ods(io.BytesIO(conteudo))
    return _linhas_csv(conteudo)


def localizar_cabecalho(linhas):
    """Procura a linha de cabeçalho e devolve (índice da linha, mapa de
    coluna -> posição). Percorre as 15 primeiras linhas porque muita planilha
    vem com título, logotipo ou linhas em branco antes da tabela."""
    melhor_indice = None
    melhor_mapa = {}

    for indice, linha in enumerate(linhas[:15]):
        rotulos = [normalizar_rotulo(celula) for celula in linha]
        mapa = {}
        for campo, aceitos in COLUNAS_RECUPERACAO.items():
            for posicao, rotulo in enumerate(rotulos):
                if not rotulo or posicao in mapa.values():
                    continue
                if rotulo in aceitos:
                    mapa[campo] = posicao
                    break
        if len(mapa) > len(melhor_mapa):
            melhor_indice, melhor_mapa = indice, mapa

    # Com menos de 3 colunas reconhecidas provavelmente não achamos a tabela.
    if melhor_indice is None or len(melhor_mapa) < 3:
        raise PlanilhaErro(
            "Não encontrei a linha de cabeçalho na planilha. Ela precisa ter "
            "uma linha com os nomes das colunas (Company, Cliente, Nota Fiscal, "
            "Emissão, Vencimento, Valor, Saldo, Juros, Correção, Outros, Total)."
        )
    return melhor_indice, melhor_mapa


def ler_planilha_recuperacao(arquivo):
    """Lê a planilha de títulos e devolve (linhas, avisos, colunas_faltando).

    Cada linha é um dicionário pronto para virar JSON: textos como string,
    datas em 'AAAA-MM-DD' (formato que o <input type=date> entende) e números
    como float. Nada é gravado no banco aqui - isso só acontece quando o
    operador salva o processo."""
    linhas_brutas = ler_linhas_brutas(arquivo)
    if not linhas_brutas:
        raise PlanilhaErro("A planilha não tem nenhuma linha.")

    indice_cabecalho, mapa = localizar_cabecalho(linhas_brutas)
    colunas_faltando = [campo for campo in COLUNAS_RECUPERACAO if campo not in mapa]

    linhas = []
    avisos = []

    for deslocamento, linha in enumerate(linhas_brutas[indice_cabecalho + 1:], start=1):
        numero_linha = indice_cabecalho + 1 + deslocamento  # como o operador vê no Excel

        # Pula linhas totalmente vazias.
        if not any(str(celula).strip() for celula in linha if celula is not None):
            continue

        registro = {}
        for campo, posicao in mapa.items():
            registro[campo] = linha[posicao] if posicao < len(linha) else None

        # Pula a linha de totais que costuma vir no rodapé da planilha
        # (sem Company/Cliente/NF, só com números).
        identificacao = " ".join(
            str(registro.get(campo) or "") for campo in ("company", "cliente", "nota_fiscal")
        ).strip()
        if not identificacao:
            continue
        if normalizar_rotulo(identificacao) in {"total", "total geral", "totais", "soma"}:
            continue

        saida = {"status": "em_analise"}
        for campo in ("company", "cliente", "nota_fiscal"):
            bruto = registro.get(campo)
            if isinstance(bruto, float) and bruto.is_integer():
                bruto = int(bruto)  # nota fiscal 1234.0 vira 1234
            saida[campo] = str(bruto).strip() if bruto is not None else ""

        for campo in CAMPOS_DATA_RECUPERACAO:
            bruto = registro.get(campo)
            convertida = valor_para_data_planilha(bruto)
            saida[campo] = convertida.isoformat() if convertida else ""
            if convertida is None and bruto not in (None, ""):
                avisos.append(
                    f"Linha {numero_linha}: não consegui entender a data de "
                    f"{'emissão' if campo == 'emissao' else 'vencimento'} "
                    f"(\"{str(bruto).strip()}\") - preencha na tela."
                )

        for campo in CAMPOS_NUMERICOS_RECUPERACAO:
            saida[campo] = valor_para_numero_planilha(registro.get(campo))

        linhas.append(saida)

    if not linhas:
        raise PlanilhaErro(
            "Encontrei o cabeçalho, mas nenhuma linha de título abaixo dele."
        )

    # Só os 12 primeiros avisos, para não virar uma parede de texto.
    if len(avisos) > 12:
        restantes = len(avisos) - 12
        avisos = avisos[:12] + [f"...e mais {restantes} aviso(s) do mesmo tipo."]

    return linhas, avisos, colunas_faltando


@app.route("/recuperacao/ler-planilha", methods=["POST"])
@login_required
def recuperacao_ler_planilha():
    """Recebe só o arquivo da planilha (via JavaScript), devolve as linhas em
    JSON. Não grava nada - quem grava é o submit do formulário."""
    arquivo = request.files.get("planilha")
    if not arquivo or not arquivo.filename:
        return jsonify({"ok": False, "erro": "Nenhum arquivo foi selecionado."}), 400

    try:
        linhas, avisos, colunas_faltando = ler_planilha_recuperacao(arquivo)
    except PlanilhaErro as erro:
        return jsonify({"ok": False, "erro": str(erro)}), 400
    except Exception:
        app.logger.exception("Falha ao ler planilha de recuperação de crédito")
        return jsonify({
            "ok": False,
            "erro": "Não consegui ler este arquivo. Confira se ele abre normalmente "
                    "no Excel e se a primeira aba é a que tem a tabela de títulos.",
        }), 400

    if colunas_faltando:
        nomes = {
            "company": "Company", "cliente": "Cliente", "nota_fiscal": "Nota Fiscal",
            "emissao": "Emissão", "vencimento": "Vencimento", "valor": "Valor",
            "saldo": "Saldo", "juros": "Juros", "correcao": "Correção",
            "outros": "Outros", "total": "Total",
        }
        avisos.insert(0, "Colunas não encontradas na planilha (ficaram em branco): "
                         + ", ".join(nomes[campo] for campo in colunas_faltando) + ".")

    return jsonify({"ok": True, "linhas": linhas, "avisos": avisos})


def preencher_titulos_recuperacao(processo, form):
    """Lê os campos repetidos da tabela de títulos (pareados pela posição,
    igual aos pedidos/verbas do trabalhista) e monta os objetos
    TituloRecuperacao."""
    colunas = {
        "company": form.getlist("rec_company"),
        "cliente": form.getlist("rec_cliente"),
        "nota_fiscal": form.getlist("rec_nota_fiscal"),
        "emissao": form.getlist("rec_emissao"),
        "vencimento": form.getlist("rec_vencimento"),
        "valor": form.getlist("rec_valor"),
        "saldo": form.getlist("rec_saldo"),
        "juros": form.getlist("rec_juros"),
        "correcao": form.getlist("rec_correcao"),
        "outros": form.getlist("rec_outros"),
        "total": form.getlist("rec_total"),
        "status": form.getlist("rec_status"),
    }
    quantidade = max((len(lista) for lista in colunas.values()), default=0)

    def pegar(campo, indice):
        lista = colunas[campo]
        return lista[indice].strip() if indice < len(lista) and lista[indice] else ""

    for indice in range(quantidade):
        textos = [pegar(campo, indice) for campo in
                  ("company", "cliente", "nota_fiscal", "emissao", "vencimento")]
        numeros = [valor_para_numero_planilha(pegar(campo, indice))
                   for campo in CAMPOS_NUMERICOS_RECUPERACAO]

        # Linha completamente em branco: ignora.
        if not any(textos) and not any(n is not None for n in numeros):
            continue

        processo.titulos_recuperacao.append(TituloRecuperacao(
            company=textos[0] or None,
            cliente=textos[1] or None,
            nota_fiscal=textos[2] or None,
            emissao=valor_para_data_planilha(textos[3]),
            vencimento=valor_para_data_planilha(textos[4]),
            valor=numeros[0],
            saldo=numeros[1],
            juros=numeros[2],
            correcao=numeros[3],
            outros=numeros[4],
            total=numeros[5],
            status=pegar("status", indice) or "em_analise",
            ordem=indice,
        ))


@app.route("/cadastro/civel-recuperacao", methods=["GET", "POST"])
@login_required
def cadastro_civel_recuperacao():
    if request.method == "GET":
        sucesso = request.args.get("sucesso") == "1"
        return render_template("cadastro_civel_recuperacao.html", sucesso=sucesso)

    form = request.form
    numero_processo = form.get("numero_processo")

    faltando = campos_obrigatorios_faltando(form)
    if faltando:
        return render_template("cadastro_civel_recuperacao.html", erro_campos_obrigatorios=faltando)

    if numero_processo_ja_existe(numero_processo):
        return render_template(
            "cadastro_civel_recuperacao.html",
            erro_numero_duplicado=numero_processo,
        )

    processo = montar_processo_base(form)
    processo.origem_cadastro = "civel_recuperacao"
    preencher_partes_e_advogados(processo, form)

    # Títulos em recuperação (linhas da tabela editável)
    preencher_titulos_recuperacao(processo, form)

    db.session.add(processo)
    quantidade_titulos = len(processo.titulos_recuperacao)
    registrar_atividade(
        "processo_criado",
        f"Cadastrou o processo cível - recuperação de crédito {numero_processo}"
        + (f" ({quantidade_titulos} título(s))" if quantidade_titulos else "")
    )
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return render_template("cadastro_civel_recuperacao.html", erro_numero_duplicado=numero_processo)

    return redirect(url_for("cadastro_civel_recuperacao", sucesso=1))


@app.route("/cadastro/trabalhista", methods=["GET", "POST"])
@login_required
def cadastro_trabalhista():
    if request.method == "GET":
        sucesso = request.args.get("sucesso") == "1"
        return render_template("cadastro_trabalhista.html", sucesso=sucesso)

    form = request.form
    numero_processo = form.get("numero_processo")

    faltando = campos_obrigatorios_faltando(form, trabalhista=True)
    if faltando:
        return render_template("cadastro_trabalhista.html", erro_campos_obrigatorios=faltando)

    if numero_processo_ja_existe(numero_processo):
        return render_template(
            "cadastro_trabalhista.html",
            erro_numero_duplicado=numero_processo,
        )

    processo = montar_processo_base(form)
    processo.origem_cadastro = "trabalhista"
    preencher_partes_e_advogados(processo, form)

    # Pedidos/verbas trabalhistas (verba + valor, pareados pela posição)
    verbas = form.getlist("pedido_verba")
    valores = form.getlist("pedido_valor")
    statuses = form.getlist("pedido_status")
    for verba, valor, status in zip(verbas, valores, statuses):
        if verba.strip():
            processo.pedidos_trabalhistas.append(
                PedidoTrabalhista(verba=verba.strip(), valor=texto_para_numero(valor), status=status or "em_analise")
            )

    db.session.add(processo)
    registrar_atividade("processo_criado", f"Cadastrou o processo trabalhista {numero_processo}")
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return render_template("cadastro_trabalhista.html", erro_numero_duplicado=numero_processo)

    return redirect(url_for("cadastro_trabalhista", sucesso=1))


@app.route("/processos/civel")
@login_required
def processos_civel():
    query = Processo.query.filter_by(origem_cadastro="civel")
    f = request.args

    numero_processo = f.get("numero_processo", "").strip()
    parte = f.get("parte", "").strip()
    juizado = f.get("juizado", "").strip()
    comarca = f.get("comarca", "").strip()
    uf = f.get("uf", "").strip()
    tipo_acao = f.get("tipo_acao", "").strip()
    centro_resultado = f.get("centro_resultado", "").strip()
    escritorio = f.get("escritorio", "").strip()
    resultado = f.get("resultado", "").strip()
    status = f.get("status", "").strip()
    risco = f.get("risco", "").strip()
    grau_instancia = f.get("grau_instancia", "").strip()
    data_distribuicao_de = f.get("data_distribuicao_de", "").strip()
    data_distribuicao_ate = f.get("data_distribuicao_ate", "").strip()

    if numero_processo:
        query = query.filter(Processo.numero_processo.ilike(f"%{numero_processo}%"))
    if parte:
        query = query.filter(Processo.partes.any(Parte.nome.ilike(f"%{parte}%")))
    if juizado:
        query = query.filter(Processo.juizado.ilike(f"%{juizado}%"))
    if comarca:
        query = query.filter(Processo.comarca.ilike(f"%{comarca}%"))
    if uf:
        query = query.filter(Processo.uf == uf)
    if tipo_acao:
        query = query.filter(Processo.tipo_acao == tipo_acao)
    if centro_resultado:
        query = query.filter(Processo.centro_resultado == centro_resultado)
    if escritorio:
        query = query.filter(Processo.escritorio == escritorio)
    if resultado:
        query = query.filter(Processo.resultado == resultado)
    if status:
        query = query.filter(Processo.status == status)
    if risco:
        query = query.filter(Processo.risco == risco)
    if grau_instancia:
        query = query.filter(Processo.grau_instancia == grau_instancia)
    if data_distribuicao_de:
        data_de = texto_para_data(data_distribuicao_de)
        if data_de:
            query = query.filter(Processo.data_distribuicao >= data_de)
    if data_distribuicao_ate:
        data_ate = texto_para_data(data_distribuicao_ate)
        if data_ate:
            query = query.filter(Processo.data_distribuicao <= data_ate)

    processos = query.order_by(Processo.data_cadastro.desc()).all()

    campos_filtro = [
        "numero_processo", "parte", "juizado", "comarca", "uf", "tipo_acao",
        "centro_resultado", "escritorio", "resultado", "status",
        "risco", "grau_instancia",
        "data_distribuicao_de", "data_distribuicao_ate",
    ]
    filtros_ativos = any(f.get(c, "").strip() for c in campos_filtro)

    return render_template(
        "processos_civel.html",
        processos=processos,
        filtros=f,
        filtros_ativos=filtros_ativos,
    )


@app.route("/processos/trabalhista")
@login_required
def processos_trabalhista():
    query = Processo.query.filter_by(origem_cadastro="trabalhista")
    f = request.args

    numero_processo = f.get("numero_processo", "").strip()
    parte = f.get("parte", "").strip()
    juizado = f.get("juizado", "").strip()
    comarca = f.get("comarca", "").strip()
    uf = f.get("uf", "").strip()
    tipo_acao = f.get("tipo_acao", "").strip()
    centro_resultado = f.get("centro_resultado", "").strip()
    escritorio = f.get("escritorio", "").strip()
    resultado = f.get("resultado", "").strip()
    sentenca = f.get("sentenca", "").strip()
    status = f.get("status", "").strip()
    risco = f.get("risco", "").strip()
    grau_instancia = f.get("grau_instancia", "").strip()
    pedido_status = f.get("pedido_status", "").strip()
    data_distribuicao_de = f.get("data_distribuicao_de", "").strip()
    data_distribuicao_ate = f.get("data_distribuicao_ate", "").strip()

    if numero_processo:
        query = query.filter(Processo.numero_processo.ilike(f"%{numero_processo}%"))
    if parte:
        query = query.filter(Processo.partes.any(Parte.nome.ilike(f"%{parte}%")))
    if juizado:
        query = query.filter(Processo.juizado.ilike(f"%{juizado}%"))
    if comarca:
        query = query.filter(Processo.comarca.ilike(f"%{comarca}%"))
    if uf:
        query = query.filter(Processo.uf == uf)
    if tipo_acao:
        query = query.filter(Processo.tipo_acao == tipo_acao)
    if centro_resultado:
        query = query.filter(Processo.centro_resultado == centro_resultado)
    if escritorio:
        query = query.filter(Processo.escritorio == escritorio)
    if resultado:
        query = query.filter(Processo.resultado == resultado)
    if sentenca:
        query = query.filter(Processo.sentenca == sentenca)
    if status:
        query = query.filter(Processo.status == status)
    if risco:
        query = query.filter(Processo.risco == risco)
    if grau_instancia:
        query = query.filter(Processo.grau_instancia == grau_instancia)
    if pedido_status:
        query = query.filter(Processo.pedidos_trabalhistas.any(PedidoTrabalhista.status == pedido_status))
    if data_distribuicao_de:
        data_de = texto_para_data(data_distribuicao_de)
        if data_de:
            query = query.filter(Processo.data_distribuicao >= data_de)
    if data_distribuicao_ate:
        data_ate = texto_para_data(data_distribuicao_ate)
        if data_ate:
            query = query.filter(Processo.data_distribuicao <= data_ate)

    processos = query.order_by(Processo.data_cadastro.desc()).all()

    campos_filtro = [
        "numero_processo", "parte", "juizado", "comarca", "uf", "tipo_acao",
        "centro_resultado", "escritorio", "resultado", "sentenca", "status",
        "risco", "grau_instancia", "pedido_status",
        "data_distribuicao_de", "data_distribuicao_ate",
    ]
    filtros_ativos = any(f.get(c, "").strip() for c in campos_filtro)

    return render_template(
        "processos_trabalhista.html",
        processos=processos,
        filtros=f,
        filtros_ativos=filtros_ativos,
    )


@app.route("/processos/civel-recuperacao")
@login_required
def processos_civel_recuperacao():
    """Listagem dos processos de Recuperação de Crédito, com a mesma
    pesquisa processual da trabalhista + os filtros próprios dos títulos
    (Company, Cliente, Nota Fiscal, vencimento e situação)."""
    query = Processo.query.filter_by(origem_cadastro="civel_recuperacao")
    f = request.args

    def texto(chave):
        return f.get(chave, "").strip()

    numero_processo = texto("numero_processo")
    parte = texto("parte")
    company = texto("company")
    cliente = texto("cliente")
    nota_fiscal = texto("nota_fiscal")
    juizado = texto("juizado")
    comarca = texto("comarca")
    uf = texto("uf")
    tipo_acao = texto("tipo_acao")
    centro_resultado = texto("centro_resultado")
    escritorio = texto("escritorio")
    resultado = texto("resultado")
    status = texto("status")
    risco = texto("risco")
    grau_instancia = texto("grau_instancia")
    titulo_status = texto("titulo_status")
    data_distribuicao_de = texto("data_distribuicao_de")
    data_distribuicao_ate = texto("data_distribuicao_ate")
    vencimento_de = texto("vencimento_de")
    vencimento_ate = texto("vencimento_ate")

    if numero_processo:
        query = query.filter(Processo.numero_processo.ilike(f"%{numero_processo}%"))
    if parte:
        query = query.filter(Processo.partes.any(Parte.nome.ilike(f"%{parte}%")))
    if company:
        query = query.filter(Processo.titulos_recuperacao.any(
            TituloRecuperacao.company.ilike(f"%{company}%")))
    if cliente:
        query = query.filter(Processo.titulos_recuperacao.any(
            TituloRecuperacao.cliente.ilike(f"%{cliente}%")))
    if nota_fiscal:
        query = query.filter(Processo.titulos_recuperacao.any(
            TituloRecuperacao.nota_fiscal.ilike(f"%{nota_fiscal}%")))
    if juizado:
        query = query.filter(Processo.juizado.ilike(f"%{juizado}%"))
    if comarca:
        query = query.filter(Processo.comarca.ilike(f"%{comarca}%"))
    if uf:
        query = query.filter(Processo.uf == uf)
    if tipo_acao:
        query = query.filter(Processo.tipo_acao == tipo_acao)
    if centro_resultado:
        query = query.filter(Processo.centro_resultado == centro_resultado)
    if escritorio:
        query = query.filter(Processo.escritorio == escritorio)
    if resultado:
        query = query.filter(Processo.resultado == resultado)
    if status:
        query = query.filter(Processo.status == status)
    if risco:
        query = query.filter(Processo.risco == risco)
    if grau_instancia:
        query = query.filter(Processo.grau_instancia == grau_instancia)
    if titulo_status:
        query = query.filter(Processo.titulos_recuperacao.any(
            TituloRecuperacao.status == titulo_status))
    if data_distribuicao_de:
        data_de = texto_para_data(data_distribuicao_de)
        if data_de:
            query = query.filter(Processo.data_distribuicao >= data_de)
    if data_distribuicao_ate:
        data_ate = texto_para_data(data_distribuicao_ate)
        if data_ate:
            query = query.filter(Processo.data_distribuicao <= data_ate)
    if vencimento_de:
        data_de = texto_para_data(vencimento_de)
        if data_de:
            query = query.filter(Processo.titulos_recuperacao.any(
                TituloRecuperacao.vencimento >= data_de))
    if vencimento_ate:
        data_ate = texto_para_data(vencimento_ate)
        if data_ate:
            query = query.filter(Processo.titulos_recuperacao.any(
                TituloRecuperacao.vencimento <= data_ate))

    processos = query.order_by(Processo.data_cadastro.desc()).all()

    campos_filtro = [
        "numero_processo", "parte", "company", "cliente", "nota_fiscal",
        "juizado", "comarca", "uf", "tipo_acao", "centro_resultado",
        "escritorio", "resultado", "status", "risco", "grau_instancia",
        "titulo_status", "data_distribuicao_de", "data_distribuicao_ate",
        "vencimento_de", "vencimento_ate",
    ]
    filtros_ativos = any(f.get(c, "").strip() for c in campos_filtro)

    # Totais do rodapé da listagem (soma de todos os títulos exibidos).
    total_titulos = sum(len(processo.titulos_recuperacao) for processo in processos)
    total_geral = sum(
        float(titulo.total or 0)
        for processo in processos
        for titulo in processo.titulos_recuperacao
    )

    return render_template(
        "processos_civel_recuperacao.html",
        processos=processos,
        filtros=f,
        filtros_ativos=filtros_ativos,
        total_titulos=total_titulos,
        total_geral=total_geral,
    )


@app.route("/agenda")
@login_required
def agenda():
    """Agenda de audiências: reúne as audiências (1, 2 e 3) de todos os
    processos cíveis, cíveis-recuperação, trabalhistas e tributários
    (estes últimos ainda são processos cíveis com tipo_acao='tributario',
    já que não existe cadastro próprio para tributário)."""
    f = request.args

    def texto(chave):
        return f.get(chave, "").strip()

    tipo = texto("tipo") or "todos"
    numero_processo = texto("numero_processo")
    parte = texto("parte")
    juizado = texto("juizado")
    comarca = texto("comarca")
    uf = texto("uf")
    tipo_audiencia = texto("tipo_audiencia")
    data_de = texto("data_de")
    data_ate = texto("data_ate")

    query = Processo.query

    if tipo == "civel":
        query = query.filter(
            Processo.origem_cadastro == "civel", Processo.tipo_acao != "tributario"
        )
    elif tipo == "civel_recuperacao":
        query = query.filter(Processo.origem_cadastro == "civel_recuperacao")
    elif tipo == "trabalhista":
        query = query.filter(Processo.origem_cadastro == "trabalhista")
    elif tipo == "tributario":
        query = query.filter(
            Processo.origem_cadastro == "civel", Processo.tipo_acao == "tributario"
        )
    # tipo == "todos" -> sem filtro de origem

    if numero_processo:
        query = query.filter(Processo.numero_processo.ilike(f"%{numero_processo}%"))
    if parte:
        query = query.filter(Processo.partes.any(Parte.nome.ilike(f"%{parte}%")))
    if juizado:
        query = query.filter(Processo.juizado.ilike(f"%{juizado}%"))
    if comarca:
        query = query.filter(Processo.comarca.ilike(f"%{comarca}%"))
    if uf:
        query = query.filter(Processo.uf == uf)

    query = query.filter(
        db.or_(
            Processo.data_audiencia_1.isnot(None),
            Processo.data_audiencia_2.isnot(None),
            Processo.data_audiencia_3.isnot(None),
        )
    )

    processos = query.all()

    # Sem filtro de data explícito, considera "próxima audiência" a partir de
    # hoje - não interessa mostrar audiências que já aconteceram.
    data_de_obj = texto_para_data(data_de) if data_de else date.today()
    data_ate_obj = texto_para_data(data_ate) if data_ate else None

    audiencias = []
    for p in processos:
        tipo_parte = "reclamante" if p.origem_cadastro == "trabalhista" else "autor"
        nomes_partes = [parte_obj.nome for parte_obj in p.partes if parte_obj.tipo == tipo_parte]
        nome_parte = ", ".join(nomes_partes) if nomes_partes else "—"

        if p.origem_cadastro == "trabalhista":
            tipo_processo = "trabalhista"
        elif p.origem_cadastro == "civel_recuperacao":
            tipo_processo = "civel_recuperacao"
        elif p.tipo_acao == "tributario":
            tipo_processo = "tributario"
        else:
            tipo_processo = "civel"

        slots = [
            (1, p.data_audiencia_1, p.audiencia_1_horario, p.audiencia_1_tipo),
            (2, p.data_audiencia_2, None, None),
            (3, p.data_audiencia_3, None, None),
        ]
        candidatos = []
        for numero_audiencia, data_aud, horario, tipo_aud in slots:
            if not data_aud:
                continue
            if data_aud < data_de_obj:
                continue
            if data_ate_obj and data_aud > data_ate_obj:
                continue
            if tipo_audiencia and tipo_aud != tipo_audiencia:
                continue
            candidatos.append((numero_audiencia, data_aud, horario, tipo_aud))

        if not candidatos:
            continue

        # De todas as audiências do processo que passaram nos filtros, só a
        # mais próxima (menor data) entra na agenda.
        numero_audiencia, data_aud, horario, tipo_aud = min(candidatos, key=lambda c: c[1])

        # Link da audiência e advogado(s) da empresa (lado "reu") - só a
        # primeira audiência tem link cadastrado (audiencia_1_link); nas
        # demais (2ª e 3ª) o campo fica vazio, já que o modelo não guarda
        # link para elas.
        link_audiencia = p.audiencia_1_link if numero_audiencia == 1 else None
        nomes_advogados = [adv.nome for adv in p.advogados if adv.lado == "reu"]
        advogados = ", ".join(nomes_advogados) if nomes_advogados else "—"

        audiencias.append({
            "processo_id": p.id,
            "numero_processo": p.numero_processo,
            "nome_parte": nome_parte,
            "tipo_processo": tipo_processo,
            "data": data_aud,
            "horario": horario,
            "tipo_audiencia": tipo_aud,
            "numero_audiencia": numero_audiencia,
            "juizado": p.juizado,
            "comarca": p.comarca,
            "uf": p.uf,
            "link": link_audiencia,
            "advogados": advogados,
        })

    audiencias.sort(key=lambda a: (a["data"], a["horario"] or ""))

    campos_filtro = [
        "numero_processo", "parte", "juizado", "comarca", "uf",
        "tipo_audiencia", "data_de", "data_ate",
    ]
    filtros_ativos = tipo != "todos" or any(f.get(c, "").strip() for c in campos_filtro)

    return render_template(
        "agenda.html",
        audiencias=audiencias,
        filtros=f,
        filtros_ativos=filtros_ativos,
        tipo_selecionado=tipo,
    )


@app.route("/processo/<int:processo_id>")
@login_required
def processo_detalhe(processo_id):
    processo = Processo.query.get_or_404(processo_id)
    return render_template("processo_detalhe.html", p=processo)


@app.route("/processo/<int:processo_id>/editar", methods=["GET", "POST"])
@login_required
def processo_editar(processo_id):
    processo = Processo.query.get_or_404(processo_id)

    if request.method == "GET":
        return render_template("processo_editar.html", p=processo)

    form = request.form
    numero_processo = form.get("numero_processo")
    numero_processo_original = processo.numero_processo

    # Bloqueia número duplicado, ignorando o próprio processo que está sendo editado
    duplicado = Processo.query.filter(
        Processo.numero_processo == numero_processo,
        Processo.id != processo_id,
    ).first()
    if duplicado:
        return render_template("processo_editar.html", p=processo, erro_numero_duplicado=numero_processo)

    preencher_campos_processo(processo, form)
    processo.numero_processo = numero_processo_original  # não pode ser alterado

    # Substitui advogados, movimentos e rateio pelo que veio no formulário -
    # mais simples e seguro do que tentar casar item a item quais foram
    # editados/removidos/adicionados. As partes (reclamante/autor e
    # reclamada/réu) NÃO são substituídas: não podem ser alteradas depois
    # do cadastro, então ficam como já estavam.
    processo.advogados = []
    processo.movimentos = []
    processo.rateio_crs = []
    preencher_partes_e_advogados(processo, form, incluir_partes=False)

    if processo.origem_cadastro == "trabalhista":
        processo.pedidos_trabalhistas = []
        verbas = form.getlist("pedido_verba")
        valores = form.getlist("pedido_valor")
        statuses = form.getlist("pedido_status")
        for verba, valor, status in zip(verbas, valores, statuses):
            if verba.strip():
                processo.pedidos_trabalhistas.append(
                    PedidoTrabalhista(verba=verba.strip(), valor=texto_para_numero(valor), status=status or "em_analise")
                )
    elif processo.origem_cadastro == "civel_recuperacao":
        # Mesma estratégia dos demais: troca a lista inteira pelo que veio
        # da tabela editável.
        processo.titulos_recuperacao = []
        preencher_titulos_recuperacao(processo, form)
    else:
        processo.pedidos_civeis = []
        for descricao in form.getlist("pedido_civel"):
            if descricao.strip():
                processo.pedidos_civeis.append(PedidoCivel(descricao=descricao.strip()))

    try:
        registrar_atividade("processo_editado", f"Editou o processo {numero_processo_original}")
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return render_template("processo_editar.html", p=processo, erro_numero_duplicado=numero_processo)

    return redirect(url_for("processo_detalhe", processo_id=processo.id))


@app.route("/panorama-juridico")
@login_required
def panorama_juridico():

    # Sem "tipo" na URL (ex.: acabou de clicar em "Panorama Jurídico" no
    # menu lateral) -> mostra a tela de escolha do tipo de processo, sem
    # gastar tempo calculando os dados do painel. Só depois que o
    # operador escolhe um dos cartões (que já vêm com ?tipo=... no link)
    # é que a lógica abaixo roda e o painel de verdade é exibido.
    #
    # "civel_trabalhista" é o único tipo que junta mais de uma origem de
    # cadastro (Cível + Trabalhista) - não existe um "todos" genérico que
    # some literalmente todos os tipos de processo do sistema.
    tipos_processo_validos = {"civel", "trabalhista", "civel_trabalhista"}
    tipo_param = request.args.get("tipo")
    if tipo_param not in tipos_processo_validos:
        return render_template("panorama_juridico.html", mostrar_menu=True)
    tipo_selecionado = tipo_param

    # Filtro por tipo de processo. "civel_trabalhista" não aplica nenhum
    # filtro extra de origem_cadastro, pois já é a junção dos dois.
    #
    # Cível - Recuperação de Crédito fica de fora deste panorama: a análise
    # de recuperação de crédito é muito diferente (não fala em risco/
    # sentença/pedidos como os outros tipos) e vai ganhar um panorama
    # próprio depois. Até lá, nenhum processo com origem_cadastro=
    # "civel_recuperacao" deve aparecer aqui - nem nas somas/contagens de
    # "Cível + Trabalhista", nem como opção selecionável no filtro.

    def query_base():
        """Ponto de partida de toda consulta nesta página - já vem com o
        filtro de tipo de processo aplicado (se houver um selecionado) e
        sempre exclui Cível - Recuperação de Crédito (ver comentário acima)."""
        query = Processo.query.filter(Processo.origem_cadastro != "civel_recuperacao")
        if tipo_selecionado != "civel_trabalhista":
            query = query.filter(Processo.origem_cadastro == tipo_selecionado)
        return query

    def contar(**filtros):
        """Conta processos que batem com os filtros dados (além do filtro
        de tipo de processo já aplicado por query_base).
        Ex: contar(origem_cadastro='civel', status='ativo')"""
        query = query_base()
        for campo, valor in filtros.items():
            query = query.filter(getattr(Processo, campo) == valor)
        return query.count()

    def somar(coluna, **filtros):
        """Soma uma coluna numérica (ex: Processo.valor_causa) para os
        processos que batem com os filtros dados (além do filtro de tipo
        de processo já aplicado). Nunca retorna None."""
        query = db.session.query(func.coalesce(func.sum(coluna), 0)).select_from(Processo)
        query = query.filter(Processo.origem_cadastro != "civel_recuperacao")
        if tipo_selecionado != "civel_trabalhista":
            query = query.filter(Processo.origem_cadastro == tipo_selecionado)
        for campo, valor in filtros.items():
            query = query.filter(getattr(Processo, campo) == valor)
        return float(query.scalar() or 0)

    # ------------------------------------------------------------------
    # 1. VISÃO GERAL (cível x trabalhista x status)
    # ------------------------------------------------------------------
    dados = {
        "total": contar(),
        "civel": {
            "total": contar(origem_cadastro="civel"),
            "ativo": contar(origem_cadastro="civel", status="ativo"),
            "arquivado": contar(origem_cadastro="civel", status="arquivado"),
            "suspenso": contar(origem_cadastro="civel", status="suspenso"),
        },
        "trabalhista": {
            "total": contar(origem_cadastro="trabalhista"),
            "ativo": contar(origem_cadastro="trabalhista", status="ativo"),
            "arquivado": contar(origem_cadastro="trabalhista", status="arquivado"),
            "suspenso": contar(origem_cadastro="trabalhista", status="suspenso"),
        },
    }

    # ------------------------------------------------------------------
    # 2. FINANCEIRO
    #    Provisionado = valor da causa dos processos ATIVOS (exposição em
    #                    aberto - o que ainda pode ser desembolsado).
    #    Gasto         = tudo que já efetivamente saiu do caixa: valor final
    #                    pago + honorários (advogado/perícia) + custas +
    #                    depósito recursal - valor do alvará (o que voltou
    #                    para o caixa).
    #    Economizado   = soma do campo "economia_gerada".
    #    Ajuste essa régua livremente se a definição da diretoria for outra.
    # ------------------------------------------------------------------
    provisionado = somar(Processo.valor_causa, status="ativo")
    valor_pago_total = somar(Processo.valor_final)
    gasto = (
        valor_pago_total
        + somar(Processo.honorarios_advogado)
        + somar(Processo.honorarios_periciais)
        + somar(Processo.custas_processuais)
        + somar(Processo.deposito_recursal)
        - somar(Processo.valor_alvara)
    )
    economizado = somar(Processo.economia_gerada)

    dados["financeiro"] = {
        "provisionado": provisionado,
        "gasto": gasto,
        "economizado": economizado,
        "valor_causa_total": somar(Processo.valor_causa),
        "valor_pago_total": valor_pago_total,
    }

    # ------------------------------------------------------------------
    # 3. RISCO - exposição em aberto (só processos ativos) por nível de risco
    # ------------------------------------------------------------------
    dados["risco"] = {
        nivel: {
            "qtd": contar(status="ativo", risco=nivel),
            "valor": somar(Processo.valor_causa, status="ativo", risco=nivel),
        }
        for nivel in ["possivel", "provavel", "remoto"]
    }
    dados["risco_sem_classificacao"] = contar(status="ativo", risco=None)

    # ------------------------------------------------------------------
    # 4. RESULTADO - processos já com desfecho, em 4 categorias:
    #    Vitórias = resultado "ganhamos" e sentença não é parcial
    #    Parciais = sentença "procedente_parcial"
    #    Acordos  = resultado "acordo"
    #    Derrotas = resultado "perdemos"
    #    Ajuste essa régua se a definição da diretoria for outra.
    # ------------------------------------------------------------------
    vitorias_qtd = query_base().filter(
        Processo.resultado == "ganhamos", Processo.sentenca != "procedente_parcial"
    ).count()
    parciais_qtd = query_base().filter(Processo.sentenca == "procedente_parcial").count()
    acordos_qtd = contar(resultado="acordo")
    derrotas_qtd = contar(resultado="perdemos")

    dados["resultado_detalhado"] = {
        "vitorias": vitorias_qtd,
        "parciais": parciais_qtd,
        "acordos": acordos_qtd,
        "derrotas": derrotas_qtd,
    }

    # Mantido para o gráfico financeiro por resultado (usa o campo "resultado" bruto)
    dados["resultado"] = {
        chave: {
            "qtd": contar(resultado=chave),
            "valor": somar(Processo.valor_final, resultado=chave),
        }
        for chave in ["ganhamos", "perdemos", "acordo"]
    }
    total_com_desfecho = vitorias_qtd + parciais_qtd + acordos_qtd + derrotas_qtd
    dados["taxa_exito"] = (
        round(100 * (vitorias_qtd + parciais_qtd) / total_com_desfecho, 1)
        if total_com_desfecho
        else None
    )

    # Quadro "Resultado dos encerrados" do painel - usa o campo SENTENÇA
    # (não o campo "resultado" acima), em 7 categorias definidas pela
    # diretoria.
    dados["resultado_sentenca"] = {
        chave: contar(sentenca=chave)
        for chave in [
            "procedente",
            "procedente_parcial",
            "acordo",
            "extinto_com_resolucao_merito",
            "extinto_sem_resolucao_merito",
            "desistencia",
            "improcedente",
        ]
    }


    # ------------------------------------------------------------------
    # 4b. MATRIZ DE RISCO - Provável/Possível/Remoto x Baixo/Médio/Alto
    #     impacto, calculado pelo valor da causa (só processos ativos).
    #     Faixas: Baixo < R$50k · Médio R$50k-200k · Alto > R$200k
    # ------------------------------------------------------------------
    def faixa_impacto(valor):
        valor = float(valor or 0)
        if valor < 50_000:
            return "baixo"
        if valor < 200_000:
            return "medio"
        return "alto"

    matriz = {nivel: {"baixo": 0, "medio": 0, "alto": 0} for nivel in ["provavel", "possivel", "remoto"]}
    ativos = query_base().filter_by(status="ativo").all()
    for p in ativos:
        if p.risco in matriz:
            matriz[p.risco][faixa_impacto(p.valor_causa)] += 1
    dados["matriz_risco"] = matriz

    # ------------------------------------------------------------------
    # 4c. ATENÇÃO DA DIRETORIA - top 5 processos ativos de risco provável,
    #     ordenados pela maior exposição financeira.
    # ------------------------------------------------------------------
    alertas_processos = (
        query_base().filter_by(status="ativo", risco="provavel")
        .order_by(Processo.valor_causa.desc())
        .limit(5)
        .all()
    )
    alertas = []
    for p in alertas_processos:
        motivo = None
        if p.pedidos_trabalhistas:
            motivo = max(p.pedidos_trabalhistas, key=lambda x: float(x.valor or 0)).verba
        elif p.pedidos_civeis:
            motivo = p.pedidos_civeis[0].descricao
        if not motivo:
            motivo = (p.tipo_acao or "").replace("_", " ").title() or "—"
        alertas.append({
            "numero_processo": p.numero_processo,
            "valor_causa": float(p.valor_causa or 0),
            "motivo": motivo,
        })
    dados["atencao_diretoria"] = alertas

    # ------------------------------------------------------------------
    # 5. TOP CENTROS DE RESULTADO por exposição em aberto
    # ------------------------------------------------------------------
    top_cr_query = db.session.query(
        Processo.centro_resultado,
        func.count(Processo.id),
        func.coalesce(func.sum(Processo.valor_causa), 0),
    ).filter(
        Processo.status == "ativo",
        Processo.centro_resultado.isnot(None),
        Processo.origem_cadastro != "civel_recuperacao",
    )
    if tipo_selecionado != "civel_trabalhista":
        top_cr_query = top_cr_query.filter(Processo.origem_cadastro == tipo_selecionado)
    top_cr = (
        top_cr_query
        .group_by(Processo.centro_resultado)
        .order_by(func.sum(Processo.valor_causa).desc())
        .limit(6)
        .all()
    )
    dados["top_centros_resultado"] = [
        {"nome": nome, "qtd": qtd, "valor": float(valor or 0)}
        for nome, qtd, valor in top_cr
    ]

    # ------------------------------------------------------------------
    # 6. EVOLUÇÃO MENSAL - novos processos distribuídos nos últimos 12 meses
    # ------------------------------------------------------------------
    hoje = date.today()
    meses_alvo = []
    ano, mes = hoje.year, hoje.month
    for _ in range(12):
        meses_alvo.append((ano, mes))
        mes -= 1
        if mes == 0:
            mes = 12
            ano -= 1
    meses_alvo.reverse()

    nomes_mes = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
                 "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

    evolucao = []
    for ano, mes in meses_alvo:
        qtd = query_base().filter(
            func.extract("year", Processo.data_distribuicao) == ano,
            func.extract("month", Processo.data_distribuicao) == mes,
        ).count()
        evolucao.append({"label": f"{nomes_mes[mes - 1]}/{str(ano)[2:]}", "qtd": qtd})
    dados["evolucao_mensal"] = evolucao

    # ------------------------------------------------------------------
    # 7. PEDIDOS E REQUERIMENTOS MAIS RECORRENTES - aponta quais itens
    #    aparecem com mais frequência e onde a tese da empresa está mais
    #    frágil. Os pedidos são diferentes conforme o tipo de processo:
    #    - Trabalhista: cada verba tem status próprio (deferido/
    #      indeferido/em análise) - "% deferido" = perda para a empresa.
    #    - Cível: o pedido em si não tem status - usamos a SENTENÇA do
    #      processo como um todo (procedente/procedente_parcial = ponto
    #      frágil) para saber o quanto aquele pedido pesa no resultado.
    # ------------------------------------------------------------------
    pedidos_trab_query = (
        db.session.query(
            PedidoTrabalhista.verba,
            func.count(PedidoTrabalhista.id),
            func.sum(case((PedidoTrabalhista.status == "deferido", 1), else_=0)),
            func.sum(case((PedidoTrabalhista.status == "indeferido", 1), else_=0)),
            func.sum(case((PedidoTrabalhista.status == "em_analise", 1), else_=0)),
        )
        .join(Processo, PedidoTrabalhista.processo_id == Processo.id)
    )
    if tipo_selecionado != "civel_trabalhista":
        pedidos_trab_query = pedidos_trab_query.filter(Processo.origem_cadastro == tipo_selecionado)
    pedidos_trab_query = (
        pedidos_trab_query.group_by(PedidoTrabalhista.verba)
        .order_by(func.count(PedidoTrabalhista.id).desc())
        .limit(8)
        .all()
    )
    dados["pedidos_recorrentes_trabalhista"] = [
        {
            "item": verba,
            "total": total,
            "deferidos": deferidos,
            "indeferidos": indeferidos,
            "em_analise": em_analise,
            "pct_deferido": round(100 * deferidos / total, 1) if total else 0,
        }
        for verba, total, deferidos, indeferidos, em_analise in pedidos_trab_query
    ]

    pedidos_civel_query = (
        db.session.query(
            PedidoCivel.descricao,
            func.count(PedidoCivel.id),
            func.sum(case((Processo.sentenca.in_(["procedente", "procedente_parcial"]), 1), else_=0)),
        )
        .join(Processo, PedidoCivel.processo_id == Processo.id)
    )
    if tipo_selecionado != "civel_trabalhista":
        pedidos_civel_query = pedidos_civel_query.filter(Processo.origem_cadastro == tipo_selecionado)
    pedidos_civel_query = (
        pedidos_civel_query.group_by(PedidoCivel.descricao)
        .order_by(func.count(PedidoCivel.id).desc())
        .limit(8)
        .all()
    )
    dados["pedidos_recorrentes_civel"] = [
        {
            "item": descricao,
            "total": total,
            "desfavoraveis": desfavoraveis,
            "pct_desfavoravel": round(100 * desfavoraveis / total, 1) if total else 0,
        }
        for descricao, total, desfavoraveis in pedidos_civel_query
    ]

    return render_template(
        "panorama_juridico.html", dados=dados, tipo_selecionado=tipo_selecionado,
        mostrar_menu=False,
    )


if __name__ == "__main__":
    app.run(debug=True)